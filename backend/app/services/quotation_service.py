from __future__ import annotations

import shutil
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..config import Settings
from ..enums import DraftStatus, MailStatus, Severity
from ..models import Mail, QuotationDraft, QuotationDraftItem, ReviewIssue
from .price_service import find_price_candidates
from .utils import sanitize_filename


def _item_text(item) -> str:
    parts = [item.product_name]
    if item.specification:
        parts.append(item.specification)
    return "\n".join(part for part in parts if part)


def _ensure_no_blocking_reviews(session: Session, mail: Mail) -> None:
    blocking = session.scalar(
        select(ReviewIssue.id).where(
            ReviewIssue.mail_id == mail.id,
            ReviewIssue.resolved.is_(False),
            ReviewIssue.severity == Severity.BLOCKING,
        )
    )
    if blocking:
        raise ValueError("검토가 필요한 필수 항목이 남아 있어 견적서를 생성할 수 없습니다.")


def create_quotation(session: Session, settings: Settings, mail: Mail) -> QuotationDraft:
    _ensure_no_blocking_reviews(session, mail)
    if not settings.quotation_template_path.exists():
        raise FileNotFoundError(f"견적서 템플릿을 찾을 수 없습니다: {settings.quotation_template_path}")

    customer_name = mail.customer_organization or mail.customer_name or "고객"
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_name = sanitize_filename(f"견적서_{customer_name}_{timestamp}.xlsx", 120)
    output_path = settings.generated_quotes_dir / output_name
    shutil.copy2(settings.quotation_template_path, output_path)

    draft = QuotationDraft(
        mail_id=mail.id,
        status=DraftStatus.DRAFT,
        file_path=str(output_path.resolve()),
        customer_name=customer_name,
        email_subject=f"[열린문디자인] 요청하신 견적서를 보내드립니다 - {customer_name}",
        email_body=(
            f"안녕하세요. {customer_name} 담당자님.\n\n"
            "요청하신 견적서를 첨부하여 보내드립니다.\n"
            "검토 후 문의사항이 있으시면 회신 부탁드립니다.\n\n"
            "감사합니다.\n열린문디자인"
        ),
    )
    session.add(draft)
    session.flush()

    total = 0
    selected: list[QuotationDraftItem] = []
    for position, item in enumerate(mail.items, start=1):
        unit_price = item.unit_price
        source: dict = {"type": "MAIL"} if unit_price is not None else {}
        if unit_price is None:
            candidates = find_price_candidates(session, item)
            exact = next((candidate for candidate in candidates if candidate.exact), None)
            if exact is None or exact.amount is None:
                raise ValueError(f"{item.product_name}의 확정 가능한 단가가 없습니다.")
            unit_price = exact.amount
            source = {
                "type": "PRICE_TABLE",
                "source_sheet": exact.source_sheet,
                "source_cell": exact.source_cell,
                "reason": exact.reason,
                "vat": exact.vat,
            }
        amount = item.amount
        if amount is None and item.quantity is not None and unit_price is not None:
            amount = int(item.quantity * unit_price)
        if amount is not None:
            total += amount
        record = QuotationDraftItem(
            draft_id=draft.id,
            position=position,
            product_name=item.product_name,
            specification=item.specification,
            quantity=item.quantity,
            unit=item.unit,
            unit_price=unit_price,
            amount=amount,
            note=item.schedule_note,
            price_source=source,
        )
        session.add(record)
        selected.append(record)

    workbook = load_workbook(output_path)
    try:
        sheet_name = settings.quotation_template_sheet
        sheet = workbook[sheet_name] if sheet_name in workbook.sheetnames else workbook.worksheets[0]
        for other in list(workbook.worksheets):
            if other != sheet:
                workbook.remove(other)
        sheet.title = "견적서"

        sheet["B3"] = f"{customer_name} 귀하"
        sheet["D4"] = mail.customer_department or "담당자 귀하"
        sheet["D5"] = datetime.now()
        sheet["D6"] = mail.delivery_place or settings.default_delivery_place
        sheet["D7"] = mail.payment_terms or settings.default_payment_terms
        sheet["D8"] = settings.default_validity
        sheet["L5"] = mail.customer_name or ""
        sheet["L6"] = mail.customer_phone or ""
        sheet["L7"] = mail.customer_email or mail.original_sender_email or ""

        for row in range(14, 24):
            for column in (3, 6, 7, 9, 12):
                sheet.cell(row, column).value = None
        for position, item in enumerate(selected[:10], start=14):
            sheet.cell(position, 3).value = _item_text(item)
            sheet.cell(position, 6).value = item.quantity
            sheet.cell(position, 7).value = item.unit_price
            sheet.cell(position, 9).value = item.amount
            sheet.cell(position, 12).value = item.note

        sheet["G24"] = "=SUM(I14:I23)"
        sheet["D10"] = "=G24"
        sheet["I10"] = "=D10"
        try:
            workbook.calculation.fullCalcOnLoad = True
            workbook.calculation.forceFullCalc = True
            workbook.calculation.calcMode = "auto"
        except AttributeError:
            pass
        workbook.save(output_path)
    finally:
        workbook.close()

    draft.total_amount = total or None
    mail.status = MailStatus.QUOTE_CREATED
    session.commit()
    session.refresh(draft)
    return draft


def approve_draft(session: Session, draft: QuotationDraft) -> QuotationDraft:
    draft.status = DraftStatus.APPROVED
    draft.approved_at = datetime.now().astimezone().replace(tzinfo=None)
    draft.mail.status = MailStatus.APPROVED
    session.commit()
    session.refresh(draft)
    return draft
